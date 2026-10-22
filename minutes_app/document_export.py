from __future__ import annotations

from io import BytesIO

from docx import Document
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
from xml.sax.saxutils import escape


def _nonempty_lines(minutes_text: str) -> list[str]:
    return [line.strip() for line in minutes_text.splitlines() if line.strip()]


def export_docx(minutes_text: str) -> bytes:
    document = Document()
    document.add_heading("Meeting Minutes", level=0)
    for line in _nonempty_lines(minutes_text):
        if len(line) < 120 and (line.startswith("#") or (line[:1].isdigit() and any(mark in line[:8] for mark in (".", ")")))):
            document.add_heading(line.lstrip("# "), level=1)
        else:
            document.add_paragraph(line)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def export_pdf(minutes_text: str) -> bytes:
    output = BytesIO()
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="MinuteHeading", parent=styles["Heading2"], textColor=colors.HexColor("#075985"), spaceBefore=12, spaceAfter=6))
    story = [Paragraph("Meeting Minutes", styles["Title"]), Spacer(1, 0.12 * inch)]
    for line in _nonempty_lines(minutes_text):
        is_heading = len(line) < 120 and (line.startswith("#") or (line[:1].isdigit() and any(mark in line[:8] for mark in (".", ")"))))
        clean = escape(line.lstrip("# ") if line.startswith("#") else line).replace("\t", "    ")
        story.append(Paragraph(clean, styles["MinuteHeading"] if is_heading else styles["BodyText"]))
        if not is_heading:
            story.append(Spacer(1, 4))
    pdf = SimpleDocTemplate(output, pagesize=letter, rightMargin=0.75 * inch, leftMargin=0.75 * inch, topMargin=0.75 * inch, bottomMargin=0.75 * inch)
    pdf.build(story)
    return output.getvalue()
