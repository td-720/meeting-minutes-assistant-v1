from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader

MAX_REFERENCE_BYTES = 8 * 1024 * 1024
MAX_TRANSCRIPT_BYTES = 2 * 1024 * 1024


def _meaningful_pdf_text(text: str) -> str:
    # Blank PDFs can yield page-number labels; those are not readable references.
    lines = []
    for line in text.splitlines():
        cleaned = re.sub(r"\s+", " ", line).strip()
        if re.fullmatch(r"(?:page\s*)?\d+(?:\s+of\s+\d+)?", cleaned, flags=re.IGNORECASE) or cleaned == "Meeting Minutes":
            continue
        if re.search(r"[A-Za-z]{2,}", cleaned):
            lines.append(cleaned)
    return "\n".join(lines).strip()


def read_transcript_text(pasted: str, uploaded_name: str | None = None, uploaded: bytes | None = None) -> str:
    if uploaded is not None:
        if len(uploaded) > MAX_TRANSCRIPT_BYTES:
            raise ValueError("Transcript files must be 2 MB or smaller.")
        if Path(uploaded_name or "").suffix.lower() != ".txt":
            raise ValueError("Transcript upload must be a .txt file.")
        text = uploaded.decode("utf-8-sig", errors="replace")
        if "\ufffd" in text:
            raise ValueError("The transcript could not be decoded as readable UTF-8 text.")
        return text.strip()
    return pasted.strip()


def read_reference(name: str, content: bytes) -> str:
    if len(content) > MAX_REFERENCE_BYTES:
        raise ValueError("Reference files must be 8 MB or smaller.")
    suffix = Path(name).suffix.lower()
    if suffix == ".txt":
        text = content.decode("utf-8-sig", errors="replace")
        if "\ufffd" in text:
            raise ValueError("The text file could not be decoded as readable UTF-8 text.")
        return text.strip()
    if suffix == ".docx":
        document = Document(BytesIO(content))
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    paragraphs.append(" | ".join(cells))
        return "\n".join(paragraphs)
    if suffix == ".pdf":
        reader = PdfReader(BytesIO(content))
        text = _meaningful_pdf_text("\n".join(page.extract_text() or "" for page in reader.pages))
        if not text:
            raise ValueError("This PDF has no extractable text. Scanned or image-only PDFs are not supported (OCR is out of scope).")
        return text
    raise ValueError("Only DOCX, text-readable PDF, and TXT reference files are supported.")
