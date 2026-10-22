from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from pypdf import PdfReader

from minutes_app.ai_service import MODEL_NAME, generate_minutes
from minutes_app.document_export import export_docx, export_pdf
from minutes_app.file_readers import read_reference, read_transcript_text
from minutes_app.local_store import add_template, clear_draft, list_templates, load_draft, save_draft


def test_txt_transcript_and_upload():
    assert read_transcript_text("  pasted  ") == "pasted"
    assert read_transcript_text("", "meeting.txt", b"hello\n") == "hello"


def test_reject_wrong_transcript_and_oversize():
    with pytest.raises(ValueError, match=".txt"):
        read_transcript_text("", "meeting.pdf", b"abc")
    with pytest.raises(ValueError, match="2 MB"):
        read_transcript_text("", "meeting.txt", b"x" * (2 * 1024 * 1024 + 1))


def test_read_docx_pdf_txt_references():
    doc = Document()
    doc.add_heading("1.0 Opening", level=1)
    doc.add_paragraph("Synthetic notes")
    data = BytesIO()
    doc.save(data)
    text = read_reference("template.docx", data.getvalue())
    assert "1.0 Opening" in text and "Synthetic notes" in text
    assert read_reference("template.txt", b"2.0 AOB") == "2.0 AOB"
    pdf_data = export_pdf("2.0 AOB\nSynthetic item")
    pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(pdf_data)).pages)
    assert "2.0 AOB" in pdf_text and "Synthetic item" in pdf_text


def test_reject_empty_or_image_only_pdf():
    empty_pdf = export_pdf("")
    with pytest.raises(ValueError, match="no extractable text"):
        read_reference("scan.pdf", empty_pdf)


def test_local_draft_and_template_store(tmp_path, monkeypatch):
    import minutes_app.local_store as store
    monkeypatch.setattr(store, "APP_DIR", tmp_path / "appdata")
    monkeypatch.setattr(store, "TEMPLATES_DIR", tmp_path / "appdata" / "templates")
    monkeypatch.setattr(store, "TEMPLATES_INDEX", tmp_path / "appdata" / "templates.json")
    monkeypatch.setattr(store, "DRAFT_PATH", tmp_path / "appdata" / "current_draft.json")
    draft = {"transcript": "synthetic transcript", "minutes": "edited minutes"}
    save_draft(draft)
    assert load_draft() == draft
    template = add_template("Sample template", ".txt", b"1.0 Opening", "1.0 Opening")
    assert list_templates()[0]["id"] == template["id"]
    clear_draft()
    assert load_draft() == {}
    assert len(list_templates()) == 1


def test_docx_export_contains_latest_edits():
    exported = Document(BytesIO(export_docx("1.0 Opening\nEdited synthetic decision")))
    text = "\n".join(p.text for p in exported.paragraphs)
    assert "Edited synthetic decision" in text


def test_pdf_export_is_valid_and_contains_latest_edits():
    output = export_pdf("1.0 Opening\nEdited synthetic decision")
    assert output.startswith(b"%PDF")
    text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(output)).pages)
    assert "Edited synthetic decision" in text


def test_ai_requires_key(monkeypatch):
    import minutes_app.ai_service as service
    monkeypatch.setattr(service, "dotenv_values", lambda _path: {})
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    with pytest.raises(RuntimeError, match="GEMINI_API_KEY"):
        generate_minutes("synthetic transcript", "synthetic notes", "1.0 Opening")


def test_ai_request_parses_structured_response(monkeypatch):
    monkeypatch.setenv("GEMINI_API_KEY", "test-only-not-a-real-key")
    captured = {}

    class FakeModels:
        def generate_content(self, **kwargs):
            captured.update(kwargs)
            return SimpleNamespace(text=json.dumps({"minutes_text": "1.0 Opening\nSynthetic decision", "unclear_attributions": ["speaker for budget statement"]}))

    class FakeClient:
        def __init__(self, api_key):
            assert api_key == "test-only-not-a-real-key"
            self.models = FakeModels()

    import minutes_app.ai_service as service
    monkeypatch.setattr(service, "dotenv_values", lambda _path: {})
    monkeypatch.setattr(service.genai, "Client", FakeClient)
    result = generate_minutes("synthetic transcript", "synthetic notes", "1.0 Opening")
    assert result == {"minutes_text": "1.0 Opening\nSynthetic decision", "unclear_attributions": ["speaker for budget statement"]}
    assert captured["model"] == MODEL_NAME
    assert captured["config"]["response_mime_type"] == "application/json"
    assert captured["config"]["response_json_schema"]["required"] == ["minutes_text", "unclear_attributions"]


def test_streamlit_app_renders_without_errors(tmp_path, monkeypatch):
    import minutes_app.local_store as store
    monkeypatch.setattr(store, "APP_DIR", tmp_path / "appdata")
    monkeypatch.setattr(store, "TEMPLATES_DIR", tmp_path / "appdata" / "templates")
    monkeypatch.setattr(store, "TEMPLATES_INDEX", tmp_path / "appdata" / "templates.json")
    monkeypatch.setattr(store, "DRAFT_PATH", tmp_path / "appdata" / "current_draft.json")
    from streamlit.testing.v1 import AppTest
    app_path = Path(__file__).resolve().parents[1] / "app.py"
    app = AppTest.from_file(str(app_path)).run()
    assert not app.exception
    assert any("Meeting Minutes Assistant" in item.value for item in app.markdown)
    markdown = [item.value for item in app.markdown]
    footer_indices = [index for index, text in enumerate(markdown) if 'class="app-footer"' in text]
    assert footer_indices and footer_indices[-1] == len(markdown) - 1
    assert "unpaid-tier" in markdown[footer_indices[-1]]
    assert "confidential" in markdown[footer_indices[-1]]
    all_markdown = "\n".join(markdown)
    assert ".app-footer" in all_markdown
    assert "font-size: .75rem" in all_markdown
    assert not any("Free AI testing:" in text for text in markdown)
